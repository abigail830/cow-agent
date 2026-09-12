from __future__ import annotations

from io import BytesIO

from docx import Document


def extract_docx_bytes(data: bytes) -> tuple[str, list[str]]:
    warnings: list[str] = []
    doc = Document(BytesIO(data))
    parts: list[str] = []

    for paragraph in doc.paragraphs:
        line = paragraph.text.strip()
        if line:
            parts.append(line)

    for table in doc.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append("\t".join(cells))
        if rows:
            parts.append("[Table]\n" + "\n".join(rows))

    if not parts:
        warnings.append("empty document")
        return "[empty document]", warnings

    if doc.tables:
        warnings.append("table layout simplified to tab-separated text")
    return "\n\n".join(parts), warnings
